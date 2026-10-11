import os
import pymupdf  # PyMuPDF
import docx
from pathlib import Path

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_data"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

# Sample JD text
SAMPLE_JD_ROLE = "Senior Full Stack Engineer (React + Python)"
SAMPLE_JD_TEXT = """Job Title: Senior Full Stack Engineer (React + Python)
Location: Bengaluru, India (Hybrid)
Experience Required: 4 - 7 years
Employment Type: Full-time

About the Role:
We are looking for a Senior Full Stack Engineer to build and scale our modern enterprise cloud applications. You will collaborate closely with product management and engineering leads to design performant web architectures, robust RESTful APIs, and responsive user interfaces.

Key Responsibilities:
- Design and develop scalable web applications using React, TypeScript, and modern state management.
- Build high-throughput REST APIs and microservices using Python (FastAPI / Django).
- Architect relational database schemas using PostgreSQL or MySQL, optimizing complex queries.
- Containerize services using Docker and manage CI/CD pipelines (GitHub Actions / GitLab CI).
- Collaborate with cloud teams to deploy applications on AWS (EC2, S3, ECS, Lambda).
- Participate in code reviews, automated unit testing (Jest, PyTest), and maintain clean code standards.

Required Qualifications & Skills:
- 4+ years of professional software engineering experience.
- Strong proficiency in modern JavaScript/TypeScript and React.
- Solid backend development experience in Python (FastAPI, Flask, or Django).
- Hands-on experience with SQL databases (PostgreSQL or MySQL) and ORMs (SQLAlchemy, Prisma).
- Proven knowledge of Docker containerization and CI/CD pipelines.
- Experience with cloud providers, preferably AWS.
- Strong problem-solving, clean code design, and cross-functional communication skills.

Nice-to-Have / Preferred:
- Experience with Redis caching and asynchronous messaging (Celery, Kafka, or RabbitMQ).
- Familiarity with AI/LLM integrations or vector databases.
- Prior experience in SaaS or B2B enterprise software development.
"""

SAMPLE_CANDIDATES = [
    {
        "id": "cand_1",
        "name": "Alex Rivera",
        "filename": "Alex_Rivera_Senior_FullStack.pdf",
        "format": "pdf",
        "content": """ALEX RIVERA
Full Stack Software Engineer | Bengaluru, KA | alex.rivera.dev@example.com | +91 98765 43210
LinkedIn: linkedin.com/in/alex-rivera-fictional | GitHub: github.com/alexrivera-fictional

PROFESSIONAL SUMMARY
Senior Full Stack Engineer with 6 years of experience architecting and delivering high-performance SaaS applications. Expert in React, TypeScript, Python (FastAPI/Django), PostgreSQL, and AWS infrastructure. Demonstrated track record of scaling user-facing applications and optimizing complex API workflows.

TECHNICAL SKILLS
- Frontend: React, TypeScript, Next.js, Redux Toolkit, Tailwind CSS, HTML5, CSS3, Jest
- Backend: Python, FastAPI, Django, RESTful APIs, GraphQL, Node.js, Celery
- Databases: PostgreSQL, MySQL, Redis, SQLAlchemy
- DevOps & Cloud: Docker, Kubernetes, AWS (EC2, S3, ECS, RDS), GitHub Actions, CI/CD
- Testing & Tools: PyTest, Jest, Git, Postman, Jira, Agile/Scrum

PROFESSIONAL EXPERIENCE
Senior Full Stack Developer | CloudScale Solutions (Bengaluru, India)
June 2022 – Present (3.5 years)
- Architected enterprise multi-tenant analytics dashboard using React, TypeScript, and FastAPI, supporting 150k monthly active users.
- Designed relational schemas and optimized queries in PostgreSQL, reducing query latency by 42%.
- Built automated CI/CD deployment pipelines using Docker and GitHub Actions to deploy containerized microservices onto AWS ECS.
- Integrated Redis caching layer which slashed peak API response times from 340ms to 45ms.
- Mentored 4 junior engineers and enforced code review standards and unit test coverage (>85% via PyTest and Jest).

Full Stack Software Engineer | Innovatech Systems (Bengaluru, India)
July 2019 – May 2022 (2.9 years)
- Developed responsive web interfaces using React and Redux, delivering 8 key product modules.
- Created robust REST APIs with Python Django and integrated third-party payment gateways.
- Containerized development environments using Docker and Docker Compose.
- Authored automated integration tests and maintained zero-downtime database migrations.

EDUCATION
Bachelor of Technology in Computer Science & Engineering
National Institute of Technology (NIT) | Graduated: 2019

PROJECT HIGHLIGHTS
- Real-Time Fleet Telemetry Platform: Built scalable FastAPI service ingesting IoT sensor data with Celery and WebSockets, visualizing metrics on dynamic React dashboards.
- Cloud Invoice Generator: Open-source Dockerized web app with PostgreSQL and AWS S3 storage.
"""
    },
    {
        "id": "cand_2",
        "name": "Priya Sharma",
        "filename": "Priya_Sharma_FullStack_Dev.docx",
        "format": "docx",
        "content": """PRIYA SHARMA
Software Engineer – Full Stack | Bengaluru, India | priya.sharma.engineer@example.com | +91 91234 56789
Portfolio: priyasharma-dev.example.com | GitHub: github.com/priyasharma-code

PROFESSIONAL SUMMARY
Results-driven Full Stack Engineer with 4.5 years of hands-on experience building web platforms using React, Python (FastAPI), and SQL databases. Strong background in REST API design, containerization with Docker, and cloud deployments on AWS.

CORE TECHNICAL SKILLS
- Languages: Python, JavaScript, TypeScript, SQL
- Frontend: React, TypeScript, HTML/CSS, Tailwind CSS, Material UI
- Backend: FastAPI, Flask, Python REST Framework, SQLAlchemy
- Databases: PostgreSQL, MySQL, SQLite, MongoDB
- Cloud & Infrastructure: Docker, AWS (S3, EC2), Git, GitHub Actions, Linux
- Testing: PyTest, React Testing Library

WORK EXPERIENCE
Full Stack Engineer | FinEdge Technologies (Bengaluru, India)
October 2021 – Present (3+ years)
- Engineered customer onboarding portal using React and TypeScript, improving user completion rate by 28%.
- Built modular microservices with FastAPI and PostgreSQL, managing secure user authentication and document verification.
- Implemented Docker containers for local and staging environments, streamlining team onboarding.
- Created automated CI/CD workflows using GitHub Actions for linting, PyTest suites, and automated builds.
- Deployed and configured staging instances on AWS EC2 and configured static asset hosting via AWS S3.

Software Developer | HexaCore Labs (Hyderabad, India)
July 2020 – September 2021 (1.2 years)
- Developed front-end components using React and styled-components for client reporting dashboards.
- Contributed to backend REST API endpoints using Python Flask and MySQL.
- Wrote unit tests and assisted in bug fixing and database maintenance tasks.

EDUCATION
B.E. in Information Technology
VTU University, Karnataka | Graduated: 2020

CERTIFICATIONS & SKILLS
- AWS Certified Cloud Practitioner
- Modern React with Redux Masterclass
"""
    },
    {
        "id": "cand_3",
        "name": "David Chen",
        "filename": "David_Chen_Frontend_Engineer.pdf",
        "format": "pdf",
        "content": """DAVID CHEN
Senior Frontend Engineer (Expanding to Full Stack) | Remote / Pune, India | david.chen.web@example.com
GitHub: github.com/davidchen-frontend | LinkedIn: linkedin.com/in/davidchen-fictional

PROFESSIONAL SUMMARY
Frontend-specialized Software Engineer with 5 years of extensive experience delivering enterprise React and TypeScript applications. Recently expanded backend skills into Python and Node.js REST services. Strong UI/UX architecture and web performance tuning expertise.

SKILLS MATRIX
- Frontend (Expert): React, TypeScript, Next.js, Redux Toolkit, Tailwind CSS, Webpack, Vite, Jest
- Backend (Intermediate): Python (Basic scripts & Flask/FastAPI prototypes), Node.js, Express
- Databases: PostgreSQL (Basics), MongoDB
- Tools & Cloud: Git, Docker (Basic usage), AWS S3, Vercel, Netlify

EXPERIENCE
Lead Frontend Developer | Omnichannel Labs (Pune, India)
February 2022 – Present (2.8 years)
- Spearheaded frontend architecture of SaaS workflow builder using React and TypeScript.
- Implemented state management using Redux Toolkit and optimized rendering bottlenecks.
- Collaborated with backend Python team on REST API contracts, Swagger specs, and error handling.
- Built automated UI component test suites achieving 80% coverage with Jest.

Frontend Developer | PixelCraft Interactive (Pune, India)
August 2019 – January 2022 (2.5 years)
- Developed responsive SPAs using React, JavaScript (ES6+), and CSS preprocessors.
- Integrated REST APIs and handled client-side caching and form validation.
- Participated in weekly code reviews and sprint planning meetings.

BACKEND TRANSITION PROJECTS
- Task Manager API: Built a Python FastAPI backend with SQLite and JWT auth to complement a React frontend.
- Simple Docker deployment: Packaged frontend and backend services into Docker containers.

EDUCATION
Bachelor of Science in Computer Science
Pune University | Graduated: 2019
"""
    },
    {
        "id": "cand_4",
        "name": "Amina Al-Mansoor",
        "filename": "Amina_Mansoor_Backend_Data.docx",
        "format": "docx",
        "content": """AMINA AL-MANSOOR
Backend & Data Platform Engineer | Bengaluru, India | amina.mansoor.dev@example.com
GitHub: github.com/aminamansoor | LinkedIn: linkedin.com/in/aminamansoor-fictional

PROFESSIONAL PROFILE
Backend Software Engineer with 4 years of experience focusing on Python API development, data pipelines, and database management. Deep expertise in Python (FastAPI/Django), PostgreSQL, Docker, and AWS. Limited hands-on frontend experience (basic React templates).

AREAS OF EXPERTISE
- Backend: Python, FastAPI, Django, REST APIs, Celery, Redis
- Databases: PostgreSQL, MySQL, Redis, Elasticsearch
- DevOps & Cloud: Docker, Docker Compose, AWS (EC2, S3, RDS, Lambda), GitHub Actions, Linux
- Frontend: Basic HTML, CSS, JavaScript (Familiar with React structure but primarily backend focused)
- Methodologies: Agile/Scrum, Test-Driven Development (TDD), PyTest

CAREER HISTORY
Backend Engineer | DataFlow Systems (Bengaluru, India)
November 2021 – Present (3 years)
- Built high-throughput data processing REST APIs using Python FastAPI and Celery workers.
- Designed database schemas, complex relational joins, and indexing strategies in PostgreSQL.
- Containerized 6 backend microservices with Docker and deployed via CI/CD to AWS EC2 and RDS.
- Implemented Redis caching and rate limiting for public-facing endpoints.
- Maintained 90%+ unit test coverage using PyTest and integrated automated SonarQube quality gates.

Junior Python Developer | Zenith IT Services (Bengaluru, India)
July 2020 – October 2021 (1.3 years)
- Developed backend scripts and web scrapers using Python and BeautifulSoup.
- Maintained Django endpoints and executed SQL queries against MySQL databases.
- Assisted DevOps engineer with basic AWS S3 bucket backups and Docker scripts.

EDUCATION
B.Tech in Computer Science
APJ Abdul Kalam Technological University | Graduated: 2020
"""
    },
    {
        "id": "cand_5",
        "name": "Marcus Vance",
        "filename": "Marcus_Vance_Junior_WebDev.pdf",
        "format": "pdf",
        "content": """MARCUS VANCE
Junior Web Developer | Hyderabad, India | marcus.vance.dev@example.com | +91 97654 32109
GitHub: github.com/marcusvance-dev

SUMMARY
Junior Developer with 1.5 years of professional experience in frontend web development. Solid foundation in HTML, CSS, JavaScript, and beginner-level React. Eager to grow full-stack abilities and learn cloud architectures.

TECHNICAL SKILLS
- Languages: JavaScript (ES6), HTML5, CSS3, Python (Beginner/Foundational)
- Frameworks: React (Beginner to Intermediate), Bootstrap, jQuery
- Tools: Git, GitHub, VS Code, Figma
- Concepts: Responsive Design, DOM Manipulation, Basic REST APIs

PROFESSIONAL EXPERIENCE
Junior Frontend Developer | BrightWeb Media (Hyderabad, India)
May 2023 – Present (1.4 years)
- Created responsive landing pages and marketing websites using HTML5, CSS3, and JavaScript.
- Updated existing client portals built with basic React components and Bootstrap styling.
- Connected frontend interfaces to third-party REST APIs using fetch and Axios.
- Fixed UI bugs across Chrome, Firefox, and Safari browsers.

Web Development Intern | CodeSphere Agency (Hyderabad, India)
January 2023 – April 2023 (4 months)
- Assisted senior developers in building WordPress and static web pages.
- Learned Git version control and collaborated on team repositories.

EDUCATION
Bachelor of Science in Information Technology
Osmania University, Hyderabad | Graduated: 2023

ADDITIONAL LEARNING
- Currently enrolled in Udemy Course: "Complete Python & FastAPI Masterclass"
- Completed: "The Complete JavaScript Bootcamp 2023"
"""
    }
]

def generate_pdf(filepath: Path, content: str):
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    rect = pymupdf.Rect(50, 50, 545, 792)
    page.insert_textbox(rect, content, fontsize=10, fontname="helv", color=(0.1, 0.1, 0.15))
    doc.save(str(filepath))
    doc.close()

def generate_docx(filepath: Path, content: str):
    doc = docx.Document()
    lines = content.strip().split("\n")
    for i, line in enumerate(lines):
        line_clean = line.strip()
        if not line_clean:
            continue
        if i == 0:
            doc.add_heading(line_clean, level=1)
        elif line_clean.isupper() and len(line_clean) < 40:
            doc.add_heading(line_clean, level=2)
        elif line_clean.startswith("- "):
            doc.add_paragraph(line_clean[2:], style='List Bullet')
        else:
            doc.add_paragraph(line_clean)
    doc.save(str(filepath))

def main():
    print(f"Generating sample candidate CVs in {SAMPLE_DIR}...")
    jd_txt_path = SAMPLE_DIR / "Sample_JD_Senior_FullStack.txt"
    jd_txt_path.write_text(SAMPLE_JD_TEXT, encoding="utf-8")
    
    jd_pdf_path = SAMPLE_DIR / "Sample_JD_Senior_FullStack.pdf"
    generate_pdf(jd_pdf_path, SAMPLE_JD_TEXT)
    
    for cand in SAMPLE_CANDIDATES:
        out_path = SAMPLE_DIR / cand["filename"]
        if cand["format"] == "pdf":
            generate_pdf(out_path, cand["content"])
            print(f"  [PDF] Created {cand['filename']}")
        elif cand["format"] == "docx":
            generate_docx(out_path, cand["content"])
            print(f"  [DOCX] Created {cand['filename']}")
    print("All sample files generated successfully!")

if __name__ == "__main__":
    main()
