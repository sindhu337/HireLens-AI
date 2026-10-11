import re
from pathlib import Path
from typing import Tuple, Dict, Any
import pymupdf  # PyMuPDF
import docx

def extract_text_from_pdf(file_path: Path) -> str:
    """Extract clean text from a PDF file using PyMuPDF."""
    text_parts = []
    doc = pymupdf.open(str(file_path))
    for page in doc:
        page_text = page.get_text("text")
        if page_text:
            text_parts.append(page_text)
    doc.close()
    raw = "\n".join(text_parts).strip()
    # Normalize hyphenated word breaks (e.g. 'develop-\nment' -> 'development')
    raw = re.sub(r'(\w+)-\n(\w+)', r'\1\2', raw)
    return raw

def extract_text_from_docx(file_path: Path) -> str:
    """Extract clean text from a DOCX file using python-docx."""
    doc = docx.Document(str(file_path))
    text_parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                text_parts.append(" | ".join(row_texts))
    return "\n".join(text_parts).strip()

def extract_text_from_file(file_path: Path) -> Tuple[str, str]:
    """
    Extracts text from PDF or DOCX.
    Returns (extracted_text, file_type).
    """
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_path), "pdf"
    elif suffix in [".docx", ".doc"]:
        return extract_text_from_docx(file_path), "docx"
    elif suffix == ".txt":
        return file_path.read_text(encoding="utf-8", errors="ignore").strip(), "txt"
    else:
        raise ValueError(f"Unsupported file format: {suffix}. Supported formats: .pdf, .docx, .txt")

def parse_resume_to_markdown(text: str, filename: str) -> str:
    """
    Parses raw resume text into structured, clean Markdown (.md) format.
    Preserves all details, metrics, and bullet statements for accurate evidence highlighting.
    """
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    if not lines:
        stem = Path(filename).stem.replace("_", " ").replace("-", " ").title()
        return f"# {stem}\n\n*No readable resume text extracted.*"

    # 1. Candidate Name
    candidate_name = lines[0].title()
    if any(kw in candidate_name.lower() for kw in ["curriculum", "resume", "cv", "page"]) or len(candidate_name) > 50:
        candidate_name = Path(filename).stem.replace("_", " ").replace("-", " ").title()

    # 2. Extract Contact Info & Links
    email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
    email = email_match.group(0) if email_match else ""

    phone_match = re.search(r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}', text)
    phone = phone_match.group(0) if phone_match else ""

    github_match = re.search(r'(?:github\.com/[a-zA-Z0-9_-]+)', text, re.IGNORECASE)
    github = github_match.group(0) if github_match else ""

    linkedin_match = re.search(r'(?:linkedin\.com/in/[a-zA-Z0-9_-]+)', text, re.IGNORECASE)
    linkedin = linkedin_match.group(0) if linkedin_match else ""

    # Subtitle line (usually line 1 or line 2 with role / location)
    subtitle = ""
    if len(lines) > 1 and ("|" in lines[1] or "developer" in lines[1].lower() or "engineer" in lines[1].lower()):
        # Filter out email/phone from subtitle if present
        sub_parts = [p.strip() for p in lines[1].split("|")]
        clean_parts = [p for p in sub_parts if "@" not in p and not re.search(r'\d{5,}', p)]
        subtitle = " | ".join(clean_parts)

    md_output = [f"# {candidate_name}"]
    if subtitle:
        md_output.append(f"**{subtitle}**")

    contact_parts = []
    if email:
        contact_parts.append(f"**Email:** {email}")
    if phone:
        contact_parts.append(f"**Phone:** {phone}")
    if contact_parts:
        md_output.append(" | ".join(contact_parts))

    links_parts = []
    if linkedin:
        links_parts.append(f"[LinkedIn](https://{linkedin})")
    if github:
        links_parts.append(f"[GitHub](https://{github})")
    if links_parts:
        md_output.append(" | ".join(links_parts))

    md_output.append("\n---\n")

    # Section Headers recognizer
    SECTION_HEADERS = [
        "PROFESSIONAL SUMMARY", "SUMMARY", "PROFESSIONAL PROFILE", "PROFILE", "ABOUT ME", "ABOUT",
        "CORE TECHNICAL SKILLS", "TECHNICAL SKILLS", "SKILLS MATRIX", "CORE COMPETENCIES", "AREAS OF EXPERTISE", "SKILLS",
        "PROFESSIONAL EXPERIENCE", "WORK EXPERIENCE", "EXPERIENCE", "EMPLOYMENT HISTORY", "WORK HISTORY",
        "EDUCATION", "ACADEMIC BACKGROUND", "ACADEMICS",
        "PROJECTS", "PERSONAL PROJECTS", "KEY PROJECTS", "PROJECT HIGHLIGHTS",
        "CERTIFICATIONS & SKILLS", "CERTIFICATIONS", "CERTIFICATIONS & LICENSES", "LICENSES & CERTIFICATIONS", "ADDITIONAL LEARNING", "ACHIEVEMENTS"
    ]

    current_section = None
    section_content = []

    def flush_section():
        if current_section and section_content:
            md_output.append(f"## {current_section.title()}")
            for item in section_content:
                md_output.append(item)
            md_output.append("\n---\n")

    # Process remaining lines (skip header lines)
    start_idx = 1
    if len(lines) > 1 and lines[1] == subtitle:
        start_idx = 2
    elif len(lines) > 2 and ("github" in lines[2].lower() or "linkedin" in lines[2].lower()):
        start_idx = 3

    for line in lines[start_idx:]:
        clean_line = line.strip()
        norm_line = re.sub(r'[*_:#]', '', clean_line).strip().upper()

        # Check if line matches a section header
        is_header = False
        matched_header = None
        for h in SECTION_HEADERS:
            if norm_line == h or norm_line == f"## {h}" or norm_line == f"### {h}":
                is_header = True
                matched_header = h
                break

        if is_header:
            flush_section()
            current_section = matched_header
            section_content = []
            continue

        if not current_section:
            # If no header yet, default to Summary
            current_section = "PROFESSIONAL SUMMARY"

        # Format line inside section
        # 1. Experience role headers typically have | or dates
        if current_section in ["PROFESSIONAL EXPERIENCE", "WORK EXPERIENCE", "EXPERIENCE", "EMPLOYMENT HISTORY", "WORK HISTORY"]:
            if ("|" in clean_line or "–" in clean_line or "-" in clean_line) and any(kw in clean_line.lower() for kw in ["present", "201", "202", "engineer", "developer", "lead", "intern", "years"]):
                if not clean_line.startswith("-") and not clean_line.startswith("•") and not clean_line.startswith("*"):
                    section_content.append(f"\n### {clean_line}\n")
                    continue

        # 2. Skills categories: e.g. "Languages: Python, JS" or "Frontend: React..."
        if current_section in ["CORE TECHNICAL SKILLS", "TECHNICAL SKILLS", "SKILLS MATRIX", "CORE COMPETENCIES", "AREAS OF EXPERTISE", "SKILLS"]:
            if ":" in clean_line and not clean_line.startswith("http"):
                cat, _, items = clean_line.partition(":")
                clean_cat = re.sub(r'^[-•*?]\s*', '', cat).strip()
                clean_items = items.strip()
                section_content.append(f"- **{clean_cat}:** {clean_items}")
                continue
            elif not clean_line.startswith("-") and not clean_line.startswith("•") and not clean_line.startswith("*"):
                section_content.append(f"- {clean_line}")
                continue

        # 3. Certifications and achievements: format each entry as clean bullet
        if current_section in ["CERTIFICATIONS & SKILLS", "CERTIFICATIONS", "CERTIFICATIONS & LICENSES", "LICENSES & CERTIFICATIONS", "ADDITIONAL LEARNING", "ACHIEVEMENTS"]:
            clean_item = re.sub(r'^[-•*?]\s*', '', clean_line).strip()
            section_content.append(f"- {clean_item}")
            continue

        # 4. Standard bullet points and paragraphs
        if clean_line.startswith("-") or clean_line.startswith("•") or clean_line.startswith("*") or clean_line.startswith("?"):
            clean_bullet = re.sub(r'^[-•*?]\s*', '', clean_line).strip()
            section_content.append(f"- {clean_bullet}")
        else:
            # If line is a wrapped continuation of previous bullet point, merge it
            if section_content and section_content[-1].startswith("- ") and not clean_line.endswith(":") and len(clean_line) > 0 and (clean_line[0].islower() or not section_content[-1].endswith(".")):
                section_content[-1] = f"{section_content[-1]} {clean_line}"
            else:
                section_content.append(clean_line)

    flush_section()

    result_md = "\n".join(md_output).strip()
    return result_md

def parse_basic_candidate_info(text: str, filename: str) -> Dict[str, Any]:
    """
    Heuristically extracts candidate name, email, phone, and structured Markdown.
    Strictly job-related; ignores protected attributes.
    """
    # 1. Email extraction
    email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
    email = email_match.group(0) if email_match else "N/A"

    # 2. Phone extraction (international and domestic formats)
    phone_match = re.search(r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}', text)
    phone = phone_match.group(0) if phone_match else "N/A"

    # 3. Name extraction: usually first non-empty line of resume or filename
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    name = "Candidate"
    if lines:
        first_line = lines[0]
        # Clean title words
        if len(first_line) < 40 and not any(kw in first_line.lower() for kw in ["curriculum", "resume", "cv", "page"]):
            name = first_line.title()
        else:
            # Fall back to filename stem
            stem = Path(filename).stem.replace("_", " ").replace("-", " ")
            name = stem.title()
    else:
        stem = Path(filename).stem.replace("_", " ").replace("-", " ")
        name = stem.title()

    # 4. Generate structured Markdown format
    markdown_text = parse_resume_to_markdown(text, filename)

    return {
        "name": name,
        "email": email,
        "phone": phone,
        "markdown": markdown_text
    }
